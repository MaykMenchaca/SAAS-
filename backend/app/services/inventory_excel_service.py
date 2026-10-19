"""Inventario completo con Excel: exportar, revisar una importación (vista previa) y aplicarla.

Flujo: se exporta el inventario (sirve de plantilla), se edita en Excel y se sube. La vista previa clasifica
cada fila en NUEVO, CAMBIO (con antes → después), SIN CAMBIOS o ERROR, con las mismas reglas que el resto del
sistema (SKU/código de barras únicos, unidades, fracciones, decimales, rangos y permisos). Solo si no hay
errores se puede aplicar, y se aplica todo o nada en una transacción. Las existencias del Excel REEMPLAZAN a las
del sistema, como un conteo físico: la diferencia queda como ajuste "Importación desde Excel".

`plan_hash` resume el plan calculado contra el estado actual; al aplicar se recalcula y, si el inventario cambió
desde la vista previa (una venta, otro ajuste), se pide revisar de nuevo en lugar de aplicar algo distinto.
"""
from __future__ import annotations

import hashlib
import json
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from marshmallow import ValidationError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from app.extensions import db
from app.models.inventory import Stock
from app.models.product import Category, Product, UnitOfMeasure
from app.schemas.product import ProductCreateSchema, ProductUpdateSchema
from app.services.inventory_count_service import InventoryCountService
from app.services.inventory_service import InventoryService
from app.services.product_service import ProductService
from app.utils.quantities import validate_quantity

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
SHEET = "Inventario"
STOCK_NOTE = "Motivo: Importación desde Excel"

# (encabezado, campo, tipo). El orden es el del archivo exportado.
COLUMNS = [
    ("SKU", "sku", "text"),
    ("Código de barras", "barcode", "text"),
    ("Nombre", "name", "text"),
    ("Categoría", "category", "text"),
    ("Unidad", "unit", "text"),
    ("Precio de venta", "sale_price", "money"),
    ("Precio de compra", "purchase_price", "money"),
    ("IVA %", "tax_rate", "money"),
    ("Existencias", "stock", "qty"),
    ("Stock mínimo", "min_stock", "qty"),
    ("Stock máximo", "max_stock", "qty"),
]
LABELS = {campo: encabezado for encabezado, campo, _ in COLUMNS}
REQUIRED_HEADERS = ("SKU",)
PRICE_FIELDS = ("sale_price", "purchase_price", "tax_rate")


class ImportFileError(ValueError):
    """El archivo no se puede leer (no es Excel, está vacío, le faltan columnas…). Responde 400."""


class ImportPlanChanged(ValueError):
    """El inventario cambió entre la vista previa y la aplicación. Responde 409."""


class ImportHasErrors(ValueError):
    """La importación tiene filas con errores: no se aplica nada. Responde 400."""

    def __init__(self, preview: dict):
        super().__init__("La importación tiene errores: corrígelos en el Excel y vuelve a subirlo.")
        self.preview = preview


def _norm(text) -> str:
    """Encabezados y nombres sin acentos, mayúsculas ni espacios extra."""
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def _can(user, permission: str) -> bool:
    return bool(user) and (user.role_name == "admin" or permission in user.permissions)


def _fmt(value) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, Decimal):
        return format(value.normalize(), "f") if value == value.to_integral_value() else str(value)
    return str(value)


def _fmt_money(value) -> str:
    return "—" if value is None else str(Decimal(value).quantize(Decimal("0.01")))


def _fmt_qty(value) -> str:
    if value is None:
        return "—"
    value = Decimal(value).normalize()
    return format(value, "f")


# ── Exportar ────────────────────────────────────────────────────────────────────────────────────────


def export_workbook(user) -> bytes:
    """Excel con el inventario activo. El costo solo va a quien puede modificar precios."""
    with_costs = _can(user, "productos:editar_precio")
    columns = [c for c in COLUMNS if with_costs or c[1] != "purchase_price"]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([c[0] for c in columns])
    header_fill = PatternFill("solid", fgColor="416180")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    rows = (
        db.session.query(Product, Stock.quantity)
        .outerjoin(Stock, Stock.product_id == Product.id)
        .filter(Product.is_active.is_(True))
        .order_by(Product.name)
        .all()
    )
    for product, quantity in rows:
        values = {
            "sku": product.sku,
            "barcode": product.barcode,
            "name": product.name,
            "category": product.category.name if product.category else None,
            "unit": product.unit.abbreviation if product.unit else None,
            "sale_price": float(product.sale_price or 0),
            "purchase_price": float(product.purchase_price or 0),
            "tax_rate": float(product.tax_rate or 0),
            "stock": float(quantity or 0),
            "min_stock": float(product.min_stock or 0),
            "max_stock": float(product.max_stock) if product.max_stock is not None else None,
        }
        ws.append([values[c[1]] for c in columns])

    widths = {"sku": 16, "barcode": 18, "name": 40, "category": 20, "unit": 9}
    for idx, (_h, campo, kind) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(campo, 15)
        if kind == "money":
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = "0.00"
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instrucciones")
    unidades = ", ".join(f"{u.abbreviation} ({u.name}{', admite decimales' if u.allows_fractions else ''})"
                         for u in UnitOfMeasure.query.order_by(UnitOfMeasure.name).all())
    for line in [
        "Cómo usar este archivo",
        "",
        "1. Edita la hoja «Inventario» y súbela en Inventario → Importar desde Excel.",
        "2. Antes de aplicar verás qué productos son nuevos, qué cambia (antes → después), qué queda igual y qué filas tienen errores.",
        "3. Solo se aplica si no hay errores, y se aplica todo junto.",
        "",
        "Columnas",
        "- SKU (obligatorio): identifica el producto. Si no existe, se crea un producto nuevo.",
        "- Para un producto nuevo también son obligatorios Nombre, Unidad y Precio de venta.",
        "- Celda vacía = no cambiar ese dato (en productos existentes).",
        "- Existencias: la cantidad física real; REEMPLAZA a la del sistema (como un conteo). La diferencia se registra como ajuste.",
        "- Unidad: abreviatura de una unidad existente. Las unidades sin decimales solo aceptan cantidades enteras.",
        "- Categoría: si no existe se crea.",
        "- Precios con máximo 2 decimales; cantidades con máximo 3; IVA entre 0 y 100.",
        "- No borres ni renombres los encabezados.",
        "",
        f"Unidades disponibles: {unidades}",
        f"Exportado el {datetime.now():%Y-%m-%d %H:%M}.",
    ]:
        info.append([line])
    info["A1"].font = Font(bold=True, size=14)
    info["A7"].font = Font(bold=True)
    info.column_dimensions["A"].width = 120

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Leer el archivo ─────────────────────────────────────────────────────────────────────────────────


@dataclass
class Row:
    number: int
    values: dict
    present: set
    errors: list = field(default_factory=list)


def read_rows(file_bytes: bytes, filename: str | None = None) -> tuple[list[Row], set]:
    """Filas del Excel con sus valores ya convertidos. Lanza ImportFileError si el archivo no sirve."""
    if not file_bytes:
        raise ImportFileError("Sube un archivo de Excel (.xlsx).")
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ImportFileError("El archivo de Excel pesa más de 5 MB.")
    if filename and not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ImportFileError("El archivo debe ser de Excel (.xlsx).")
    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ImportFileError("No se pudo leer el archivo: debe ser un Excel (.xlsx) válido.") from exc

    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise ImportFileError("El Excel está vacío: falta la fila de encabezados.")
    by_norm = {_norm(h): campo for h, campo, _ in COLUMNS}
    positions: dict[str, int] = {}
    for idx, h in enumerate(header):
        campo = by_norm.get(_norm(h))
        if campo and campo not in positions:
            positions[campo] = idx
    missing = [h for h in REQUIRED_HEADERS if by_norm[_norm(h)] not in positions]
    if missing:
        raise ImportFileError(f"Al Excel le falta la columna obligatoria: {', '.join(missing)}.")

    kinds = {campo: kind for _h, campo, kind in COLUMNS}
    rows: list[Row] = []
    for number, raw in enumerate(rows_iter, start=2):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportFileError(f"El Excel tiene más de {MAX_ROWS} filas; divídelo en varios archivos.")
        values, present, errors = {}, set(), []
        for campo, idx in positions.items():
            cell = raw[idx] if idx < len(raw) else None
            if cell is None or (isinstance(cell, str) and cell.strip() == ""):
                continue
            present.add(campo)
            kind = kinds[campo]
            if kind == "text":
                values[campo] = str(cell).strip() if not isinstance(cell, float) or not cell.is_integer() else str(int(cell))
            else:
                try:
                    dec = Decimal(str(cell).strip().replace(",", ""))
                    if not dec.is_finite():
                        raise InvalidOperation
                    values[campo] = dec
                except (InvalidOperation, ValueError):
                    errors.append(f"{LABELS[campo]}: «{cell}» no es un número.")
        rows.append(Row(number, values, present, errors))
    if not rows:
        raise ImportFileError("El Excel no tiene filas de productos debajo de los encabezados.")
    return rows, set(positions)


# ── Plan (vista previa) ─────────────────────────────────────────────────────────────────────────────


def _schema_errors(err: ValidationError) -> list[str]:
    out = []
    for campo, msgs in err.messages.items():
        label = LABELS.get(campo, {"unit_id": "Unidad", "category_id": "Categoría"}.get(campo, campo))
        for msg in msgs if isinstance(msgs, list) else [msgs]:
            out.append(f"{label}: {msg}")
    return out


def build_plan(rows: list[Row], user) -> dict:
    can_create = _can(user, "productos:crear")
    can_edit = _can(user, "productos:editar")
    can_price = _can(user, "productos:editar_precio")
    can_stock = _can(user, "inventario:ajustar")

    units = {}
    for u in UnitOfMeasure.query.all():
        units[_norm(u.abbreviation)] = u
        units.setdefault(_norm(u.name), u)
    categories = {_norm(c.name): c for c in Category.query.all()}
    new_categories: set[str] = set()

    sku_seen: dict[str, int] = {}
    barcode_seen: dict[str, int] = {}
    result_rows, ops = [], []

    for row in rows:
        v = row.values
        errors = list(row.errors)
        sku = (v.get("sku") or "").strip()
        out = {"fila": row.number, "sku": sku, "nombre": v.get("name"), "estado": "igual", "cambios": [], "errores": errors}
        result_rows.append(out)
        if not sku:
            errors.append("Falta el SKU.")
            out["estado"] = "error"
            continue

        key = sku.lower()
        if key in sku_seen:
            errors.append(f"SKU repetido en el Excel (también en la fila {sku_seen[key]}).")
        sku_seen.setdefault(key, row.number)
        if v.get("barcode"):
            bkey = v["barcode"].strip().lower()
            if bkey in barcode_seen:
                errors.append(f"Código de barras repetido en el Excel (también en la fila {barcode_seen[bkey]}).")
            barcode_seen.setdefault(bkey, row.number)

        product = Product.query.filter(func.lower(func.trim(Product.sku)) == key).first()
        if product is not None and not product.is_active:
            errors.append("El producto está desactivado: reactívalo antes de importarlo.")
            out["estado"] = "error"
            continue
        out["nombre"] = v.get("name") or (product.name if product else None)

        # Unidad y categoría por nombre
        unit = None
        if "unit" in v:
            unit = units.get(_norm(v["unit"]))
            if unit is None:
                errors.append(f"Unidad «{v['unit']}» no existe. Usa la abreviatura de una unidad registrada.")
        category, category_label = None, None
        if "category" in v:
            category = categories.get(_norm(v["category"]))
            if category is None:
                category_label = v["category"]
                if not can_create:
                    errors.append("Sin permiso para crear la categoría nueva.")
            elif not category.is_active:
                errors.append(f"La categoría «{category.name}» está desactivada.")

        data = {k: v[k] for k in ("sku", "barcode", "name", "sale_price", "purchase_price", "tax_rate", "min_stock", "max_stock") if k in v}
        if unit is not None:
            data["unit_id"] = unit.id
        if category is not None:
            data["category_id"] = category.id
        stock_new = v.get("stock")
        if stock_new is not None and stock_new < 0:
            errors.append("Existencias: no puede ser negativa.")

        if product is None:
            # ── Producto nuevo
            out["estado"] = "nuevo"
            if not can_create:
                errors.append("Sin permiso para crear productos.")
            for req, label in (("name", "Nombre"), ("unit", "Unidad"), ("sale_price", "Precio de venta")):
                if req not in v:
                    errors.append(f"Falta {label} (obligatorio para un producto nuevo).")
            payload = {k: (str(x) if isinstance(x, Decimal) else x) for k, x in data.items()}
            payload.setdefault("unit_id", unit.id if unit else 1)
            try:
                loaded = ProductCreateSchema().load(payload)
            except ValidationError as err:
                errors.extend(_schema_errors(err))
                loaded = None
            try:
                ProductService._validate_unique_product(sku=sku, barcode=v.get("barcode"))
            except ValidationError as err:
                errors.extend(_schema_errors(err))
            if unit is not None and stock_new:
                try:
                    validate_quantity(stock_new, unit, v.get("name") or sku)
                except ValueError as exc:
                    errors.append(f"Existencias: {exc}")
            if category_label:
                new_categories.add(category_label.strip())
            for campo in ("barcode", "name", "category", "unit", "sale_price", "purchase_price", "tax_rate", "stock", "min_stock", "max_stock"):
                if campo in v:
                    despues = _fmt_money(v[campo]) if campo in PRICE_FIELDS else (_fmt_qty(v[campo]) if campo in ("stock", "min_stock", "max_stock") else str(v[campo]))
                    if campo == "category" and category_label:
                        despues = f"{category_label} (categoría nueva)"
                    out["cambios"].append({"campo": LABELS[campo], "antes": "—", "despues": despues})
            if not errors and loaded is not None:
                ops.append({"tipo": "crear", "fila": row.number, "data": loaded, "categoria_nueva": category_label,
                            "stock": stock_new or Decimal("0")})
        else:
            # ── Producto existente: solo lo que cambia
            payload = {k: (str(x) if isinstance(x, Decimal) else x) for k, x in data.items()}
            try:
                loaded = ProductUpdateSchema().load(payload)
            except ValidationError as err:
                errors.extend(_schema_errors(err))
                loaded = {}
            try:
                ProductService._validate_unique_product(sku=loaded.get("sku"), barcode=loaded.get("barcode"), exclude_product_id=product.id)
            except ValidationError as err:
                errors.extend(_schema_errors(err))
            changes = {}
            for campo, valor in loaded.items():
                actual = getattr(product, campo)
                if campo == "sku" and str(actual).strip().lower() == str(valor).strip().lower():
                    continue
                if isinstance(actual, Decimal) or isinstance(valor, Decimal):
                    if actual is not None and Decimal(str(valor)) == Decimal(actual):
                        continue
                elif (actual or None) == (valor or None):
                    continue
                changes[campo] = valor
            if category_label:
                changes["category_new"] = category_label
                new_categories.add(category_label.strip())
            target_unit = unit or product.unit
            stock_row = Stock.query.filter_by(product_id=product.id).first()
            stock_actual = stock_row.quantity if stock_row else Decimal("0")
            stock_change = stock_new is not None and Decimal(stock_new) != Decimal(stock_actual)
            if stock_new is not None:
                try:
                    if stock_new > 0:
                        validate_quantity(stock_new, target_unit, product.name)
                except ValueError as exc:
                    errors.append(f"Existencias: {exc}")
            if "unit_id" in changes and not unit.allows_fractions:
                cantidad = stock_new if stock_new is not None else stock_actual
                if cantidad % 1:
                    errors.append("No se puede cambiar a una unidad sin decimales con existencias fraccionarias.")

            # Permisos por tipo de cambio
            if any(c in PRICE_FIELDS for c in changes) and not can_price:
                errors.append("Sin permiso para modificar precios o costos.")
            if any(c not in PRICE_FIELDS for c in changes) and not can_edit:
                errors.append("Sin permiso para editar productos.")
            if stock_change and not can_stock:
                errors.append("Sin permiso para ajustar existencias.")

            for campo, valor in changes.items():
                if campo == "category_new":
                    out["cambios"].append({"campo": "Categoría", "antes": product.category.name if product.category else "—",
                                           "despues": f"{valor} (categoría nueva)"})
                    continue
                actual = getattr(product, campo)
                if campo in PRICE_FIELDS:
                    antes, despues = _fmt_money(actual), _fmt_money(valor)
                elif campo in ("min_stock", "max_stock"):
                    antes, despues = _fmt_qty(actual), _fmt_qty(valor)
                elif campo == "unit_id":
                    antes, despues = product.unit.abbreviation, unit.abbreviation
                elif campo == "category_id":
                    antes, despues = (product.category.name if product.category else "—"), category.name
                else:
                    antes, despues = _fmt(actual), _fmt(valor)
                label = {"unit_id": "Unidad", "category_id": "Categoría"}.get(campo, LABELS.get(campo, campo))
                out["cambios"].append({"campo": label, "antes": antes, "despues": despues})
            if stock_change:
                out["cambios"].append({"campo": "Existencias", "antes": _fmt_qty(stock_actual), "despues": _fmt_qty(stock_new)})

            if out["cambios"]:
                out["estado"] = "cambio"
            if not errors and out["cambios"]:
                ops.append({"tipo": "editar", "fila": row.number, "product_id": product.id,
                            "data": {k: x for k, x in changes.items() if k != "category_new"},
                            "categoria_nueva": changes.get("category_new"),
                            "stock": stock_new if stock_change else None, "stock_antes": stock_actual if stock_change else None})
        if errors:
            out["estado"] = "error"

    resumen = {"filas": len(result_rows),
               "nuevos": sum(1 for r in result_rows if r["estado"] == "nuevo"),
               "cambios": sum(1 for r in result_rows if r["estado"] == "cambio"),
               "sin_cambios": sum(1 for r in result_rows if r["estado"] == "igual"),
               "errores": sum(1 for r in result_rows if r["estado"] == "error")}
    plan_hash = hashlib.sha256(json.dumps([{**r, "errores": len(r["errores"])} for r in result_rows],
                                          sort_keys=True, default=str).encode()).hexdigest()[:32]
    return {"resumen": resumen, "filas": result_rows, "categorias_nuevas": sorted(new_categories),
            "plan_hash": plan_hash, "_ops": ops}


def preview(file_bytes: bytes, filename: str | None, user) -> dict:
    rows, _ = read_rows(file_bytes, filename)
    plan = build_plan(rows, user)
    plan.pop("_ops")
    return plan


# ── Aplicar ─────────────────────────────────────────────────────────────────────────────────────────


def apply(file_bytes: bytes, filename: str | None, user, plan_hash: str | None) -> dict:
    """Aplica la importación en una sola transacción. Recalcula el plan y exige que coincida con la vista previa."""
    from app.services.audit_service import log_action

    rows, _ = read_rows(file_bytes, filename)
    plan = build_plan(rows, user)
    ops = plan.pop("_ops")
    if plan["resumen"]["errores"]:
        raise ImportHasErrors(plan)
    if not plan_hash or plan_hash != plan["plan_hash"]:
        raise ImportPlanChanged(
            "El inventario cambió desde la vista previa (por ejemplo, hubo una venta o un ajuste). "
            "Revisa la vista previa otra vez antes de aplicar."
        )

    try:
        categorias = {}
        for nombre in plan["categorias_nuevas"]:
            cat = Category(name=nombre)
            db.session.add(cat)
            db.session.flush()
            categorias[_norm(nombre)] = cat
            log_action(action="crear", entity_type="category", entity_id=cat.id,
                       details={"name": nombre, "origen": "importacion_excel"}, user_id=user.id)

        product_ids = [op["product_id"] for op in ops if op["tipo"] == "editar" and op["stock"] is not None]
        if product_ids:
            InventoryService.lock_stock(product_ids)

        for op in ops:
            data = dict(op["data"])
            if op["categoria_nueva"]:
                data["category_id"] = categorias[_norm(op["categoria_nueva"])].id
            if op["tipo"] == "crear":
                data.pop("avg_cost", None)
                product = Product(**data)
                product.avg_cost = product.purchase_price or 0  # como el alta normal (H-26)
                db.session.add(product)
                db.session.flush()
                db.session.add(Stock(product_id=product.id, quantity=0))
                db.session.flush()
                log_action(action="crear", entity_type="product", entity_id=product.id,
                           details={"sku": product.sku, "name": product.name, "origen": "importacion_excel"},
                           new_value={"sku": product.sku, "name": product.name, "sale_price": str(product.sale_price)},
                           user_id=user.id)
                if op["stock"]:
                    InventoryService.move_stock(product_id=product.id, quantity=op["stock"], movement_type="ajuste",
                                                reference_type="ajuste_manual", reference_id=0, user_id=user.id,
                                                notes=STOCK_NOTE)
            else:
                product = db.session.get(Product, op["product_id"])
                if data:
                    previous = {k: (str(getattr(product, k)) if getattr(product, k) is not None else None) for k in data}
                    for k, x in data.items():
                        setattr(product, k, x)
                    if data.get("purchase_price") and not (product.avg_cost or 0):
                        product.avg_cost = data["purchase_price"]  # mismo criterio que la edición normal (RF-07)
                    log_action(action="editar", entity_type="product", entity_id=product.id,
                               details={"sku": product.sku, "name": product.name, "origen": "importacion_excel"},
                               previous_value=previous,
                               new_value={k: (str(getattr(product, k)) if getattr(product, k) is not None else None) for k in data},
                               user_id=user.id)
                if op["stock"] is not None:
                    # Igual que un ajuste manual: anula la captura de un conteo abierto (RF-32)
                    InventoryCountService.reset_capture_after_adjust(product.id)
                    InventoryService.move_stock(product_id=product.id, quantity=op["stock"], movement_type="ajuste",
                                                reference_type="ajuste_manual", reference_id=0, user_id=user.id,
                                                notes=STOCK_NOTE)

        log_action(action="importar_inventario", entity_type="inventory", details={
            "archivo": filename, "nuevos": plan["resumen"]["nuevos"], "cambios": plan["resumen"]["cambios"],
            "sin_cambios": plan["resumen"]["sin_cambios"], "categorias_nuevas": plan["categorias_nuevas"],
        }, user_id=user.id)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise
    return {"resumen": plan["resumen"], "categorias_nuevas": plan["categorias_nuevas"]}
