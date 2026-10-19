"""Inventario con Excel: exportar, vista previa de la importación (nuevos / cambios / sin cambios / errores) y aplicar."""
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.models.audit import AuditLog
from app.models.inventory import Stock, StockMovement
from app.models.product import Category, Product, UnitOfMeasure
from app.models.user import Role, User


def _h(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def catalogo(db):
    pz = UnitOfMeasure(name="Pieza", abbreviation="pz", allows_fractions=False)
    kg = UnitOfMeasure(name="Kilogramo", abbreviation="kg", allows_fractions=True)
    cat = Category(name="Herramientas")
    db.session.add_all([pz, kg, cat])
    db.session.flush()
    martillo = Product(sku="MAR-1", barcode="7501000000011", name="Martillo", unit_id=pz.id, category_id=cat.id,
                       sale_price=Decimal("100.00"), purchase_price=Decimal("60.00"), avg_cost=Decimal("60.00"),
                       tax_rate=Decimal("16.00"), min_stock=Decimal("5"), max_stock=Decimal("50"))
    clavo = Product(sku="CLV-KG", name="Clavo a granel", unit_id=kg.id, category_id=cat.id,
                    sale_price=Decimal("40.00"), purchase_price=Decimal("25.00"), avg_cost=Decimal("25.00"),
                    tax_rate=Decimal("16.00"))
    db.session.add_all([martillo, clavo])
    db.session.flush()
    db.session.add_all([Stock(product_id=martillo.id, quantity=Decimal("20")), Stock(product_id=clavo.id, quantity=Decimal("12.5"))])
    db.session.commit()
    return {"martillo": martillo, "clavo": clavo, "cat": cat}


def _exportar(client, token):
    resp = client.get("/api/inventory/export", headers=_h(token))
    assert resp.status_code == 200, resp.data[:200]
    return load_workbook(BytesIO(resp.data))


def _hoja_a_filas(wb):
    ws = wb["Inventario"]
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    return filas[0], filas[1:]


def _xlsx(encabezados, filas) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(encabezados)
    for fila in filas:
        ws.append(fila)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _subir(client, token, url, archivo, **form):
    data = {"archivo": (archivo, "inventario.xlsx"), **form}
    return client.post(url, headers=_h(token), data=data, content_type="multipart/form-data")


def _preview(client, token, archivo):
    return _subir(client, token, "/api/inventory/import/preview", archivo)


def _por_sku(preview):
    return {f["sku"]: f for f in preview["filas"]}


# ── Exportar ────────────────────────────────────────────────────────────────────────────────────────


def test_exportar_trae_productos_existencias_y_hoja_de_instrucciones(client, admin_token, catalogo):
    wb = _exportar(client, admin_token)
    assert "Instrucciones" in wb.sheetnames
    encabezados, filas = _hoja_a_filas(wb)
    assert encabezados[:3] == ["SKU", "Código de barras", "Nombre"]
    assert "Existencias" in encabezados and "Precio de compra" in encabezados
    por_sku = {f[0]: dict(zip(encabezados, f)) for f in filas}
    assert por_sku["MAR-1"]["Existencias"] == 20
    assert por_sku["CLV-KG"]["Existencias"] == 12.5
    assert por_sku["MAR-1"]["Unidad"] == "pz"
    assert por_sku["MAR-1"]["Categoría"] == "Herramientas"


def test_exportar_sin_permiso_de_precios_no_incluye_costos(client, regular_user_token, catalogo):
    encabezados, _ = _hoja_a_filas(_exportar(client, regular_user_token))
    assert "Precio de compra" not in encabezados
    assert "Precio de venta" in encabezados


# ── Vista previa ────────────────────────────────────────────────────────────────────────────────────


def test_reimportar_lo_exportado_no_tiene_cambios(client, admin_token, catalogo):
    encabezados, filas = _hoja_a_filas(_exportar(client, admin_token))
    resp = _preview(client, admin_token, _xlsx(encabezados, filas))
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert data["resumen"] == {"filas": 2, "nuevos": 0, "cambios": 0, "sin_cambios": 2, "errores": 0}
    assert data["plan_hash"]


def test_vista_previa_separa_nuevos_cambios_iguales_y_errores_sin_aplicar_nada(client, admin_token, db, catalogo):
    encabezados, filas = _hoja_a_filas(_exportar(client, admin_token))
    col = {n: i for i, n in enumerate(encabezados)}
    filas = [list(f) for f in filas]
    mar = next(f for f in filas if f[col["SKU"]] == "MAR-1")
    mar[col["Precio de venta"]] = 110
    mar[col["Existencias"]] = 30
    nuevo = [None] * len(encabezados)
    for k, v in {"SKU": "DES-1", "Nombre": "Desarmador", "Unidad": "pz", "Categoría": "Desarmadores",
                 "Precio de venta": 55.5, "Precio de compra": 30, "IVA %": 16, "Existencias": 10}.items():
        nuevo[col[k]] = v
    mala_unidad = [None] * len(encabezados)
    for k, v in {"SKU": "X-1", "Nombre": "Raro", "Unidad": "litro", "Precio de venta": 10}.items():
        mala_unidad[col[k]] = v
    media_pieza = [None] * len(encabezados)
    for k, v in {"SKU": "X-2", "Nombre": "Media", "Unidad": "pz", "Precio de venta": 10, "Existencias": 2.5}.items():
        media_pieza[col[k]] = v
    duplicado = list(nuevo)
    duplicado[col["SKU"]] = " des-1 "
    resp = _preview(client, admin_token, _xlsx(encabezados, filas + [nuevo, mala_unidad, media_pieza, duplicado]))
    assert resp.status_code == 200, resp.get_json()
    data = resp.get_json()
    assert data["resumen"] == {"filas": 6, "nuevos": 1, "cambios": 1, "sin_cambios": 1, "errores": 3}
    filas_r = _por_sku(data)
    cambios = {c["campo"]: (c["antes"], c["despues"]) for c in filas_r["MAR-1"]["cambios"]}
    assert cambios["Precio de venta"] == ("100.00", "110.00")
    assert cambios["Existencias"] == ("20", "30")
    assert filas_r["DES-1"]["estado"] == "nuevo"
    assert any("categoría nueva" in c["despues"].lower() for c in filas_r["DES-1"]["cambios"] if c["campo"] == "Categoría")
    assert any("unidad" in e.lower() for e in filas_r["X-1"]["errores"])
    assert any("fraccion" in e.lower() for e in filas_r["X-2"]["errores"])
    assert any("repetido" in e.lower() for e in filas_r["des-1"]["errores"])  # mismo SKU con espacios y minúsculas
    # La vista previa no cambia nada
    db.session.expire_all()
    assert db.session.get(Product, catalogo["martillo"].id).sale_price == Decimal("100.00")
    assert Product.query.filter_by(sku="DES-1").first() is None


# ── Aplicar ─────────────────────────────────────────────────────────────────────────────────────────


def _archivo_valido(client, token):
    encabezados, filas = _hoja_a_filas(_exportar(client, token))
    col = {n: i for i, n in enumerate(encabezados)}
    filas = [list(f) for f in filas]
    mar = next(f for f in filas if f[col["SKU"]] == "MAR-1")
    mar[col["Precio de venta"]] = 110
    mar[col["Existencias"]] = 30
    clv = next(f for f in filas if f[col["SKU"]] == "CLV-KG")
    clv[col["Existencias"]] = 10.25
    nuevo = [None] * len(encabezados)
    for k, v in {"SKU": "DES-1", "Nombre": "Desarmador", "Unidad": "pz", "Categoría": "Desarmadores",
                 "Precio de venta": 55.5, "Precio de compra": 30, "IVA %": 16, "Existencias": 10}.items():
        nuevo[col[k]] = v
    return encabezados, filas + [nuevo]


def test_aplicar_importacion_valida(client, admin_token, db, catalogo):
    encabezados, filas = _archivo_valido(client, admin_token)
    preview = _preview(client, admin_token, _xlsx(encabezados, filas)).get_json()
    resp = _subir(client, admin_token, "/api/inventory/import/apply", _xlsx(encabezados, filas), plan_hash=preview["plan_hash"])
    assert resp.status_code == 200, resp.get_json()
    assert resp.get_json()["resumen"]["nuevos"] == 1
    db.session.expire_all()
    mar = db.session.get(Product, catalogo["martillo"].id)
    assert mar.sale_price == Decimal("110.00")
    assert Stock.query.filter_by(product_id=mar.id).one().quantity == Decimal("30")
    assert Stock.query.filter_by(product_id=catalogo["clavo"].id).one().quantity == Decimal("10.25")
    nuevo = Product.query.filter_by(sku="DES-1").one()
    assert nuevo.avg_cost == Decimal("30.00") and nuevo.category.name == "Desarmadores"
    assert Stock.query.filter_by(product_id=nuevo.id).one().quantity == Decimal("10")
    mov = StockMovement.query.filter_by(product_id=mar.id, movement_type="ajuste").order_by(StockMovement.id.desc()).first()
    assert mov.quantity == Decimal("10") and "Excel" in (mov.notes or "")
    assert AuditLog.query.filter_by(action="importar_inventario").count() == 1


def test_con_errores_no_se_aplica_nada(client, admin_token, db, catalogo):
    encabezados, filas = _archivo_valido(client, admin_token)
    col = {n: i for i, n in enumerate(encabezados)}
    filas[0][col["Precio de venta"]] = -5
    resp = _subir(client, admin_token, "/api/inventory/import/apply", _xlsx(encabezados, filas), plan_hash="x")
    assert resp.status_code == 400
    db.session.expire_all()
    assert Product.query.filter_by(sku="DES-1").first() is None
    assert db.session.get(Product, catalogo["martillo"].id).sale_price == Decimal("100.00")


def test_si_el_inventario_cambio_desde_la_vista_previa_pide_revisar_de_nuevo(client, admin_token, db, catalogo):
    encabezados, filas = _archivo_valido(client, admin_token)
    preview = _preview(client, admin_token, _xlsx(encabezados, filas)).get_json()
    # Mientras tanto se vende o se ajusta: el stock ya no es el de la vista previa
    Stock.query.filter_by(product_id=catalogo["martillo"].id).one().quantity = Decimal("18")
    db.session.commit()
    resp = _subir(client, admin_token, "/api/inventory/import/apply", _xlsx(encabezados, filas), plan_hash=preview["plan_hash"])
    assert resp.status_code == 409
    db.session.expire_all()
    assert Product.query.filter_by(sku="DES-1").first() is None


# ── Permisos y archivos inválidos ───────────────────────────────────────────────────────────────────


@pytest.fixture
def ajustador_token(client, db, seed_permissions):
    """Puede ajustar inventario y editar productos, pero no crear productos ni cambiar precios."""
    rol = Role(name="ajustador", description="Ajustador")
    rol.permissions.extend([seed_permissions[c] for c in ("inventario:ver", "inventario:ajustar", "productos:ver", "productos:editar")])
    db.session.add(rol)
    db.session.flush()
    u = User(username="ajustador", email="aj@test.com", full_name="Ajustador", role_id=rol.id)
    u.set_password("ajustador123")
    db.session.add(u)
    db.session.commit()
    return client.post("/api/auth/login", json={"username": "ajustador", "password": "ajustador123"}).get_json()["access_token"]


def test_sin_permiso_de_crear_ni_de_precios_esas_filas_son_errores(client, admin_token, ajustador_token, catalogo):
    encabezados, filas = _archivo_valido(client, admin_token)
    data = _preview(client, ajustador_token, _xlsx(encabezados, filas)).get_json()
    filas_r = _por_sku(data)
    assert any("permiso" in e.lower() for e in filas_r["DES-1"]["errores"])
    assert any("permiso" in e.lower() for e in filas_r["MAR-1"]["errores"])  # cambia el precio
    assert filas_r["CLV-KG"]["estado"] == "cambio"  # solo existencias: sí puede


def test_cajero_no_puede_importar(client, regular_user_token, catalogo):
    resp = _preview(client, regular_user_token, _xlsx(["SKU"], [["MAR-1"]]))
    assert resp.status_code == 403


@pytest.mark.parametrize("contenido, nombre", [(b"no soy excel", "inventario.xlsx"), (b"a,b\n1,2", "inventario.csv")])
def test_archivo_que_no_es_excel_se_rechaza(client, admin_token, catalogo, contenido, nombre):
    resp = client.post("/api/inventory/import/preview", headers=_h(admin_token),
                       data={"archivo": (BytesIO(contenido), nombre)}, content_type="multipart/form-data")
    assert resp.status_code == 400
    assert "excel" in resp.get_json()["error"].lower()


def test_faltan_columnas_obligatorias(client, admin_token, catalogo):
    resp = _preview(client, admin_token, _xlsx(["Nombre", "Precio de venta"], [["X", 10]]))
    assert resp.status_code == 400
    assert "SKU" in resp.get_json()["error"]
