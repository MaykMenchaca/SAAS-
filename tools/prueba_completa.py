r"""Prueba completa del POS: recorre TODAS las funciones de la app contra una base de prueba aislada.

Crea un esquema nuevo en pos_test_db (nunca pos_db), lo migra y siembra, y prueba la API real como lo haría
la app: sesión y usuarios, catálogo, compras, inventario y conteo, caja, ventas (efectivo, mixta, a granel,
idempotencia), devoluciones, cancelación, corte de caja, reportes, configuración, auditoría, permisos por rol
y respaldos. Recalcula los montos con Decimal y compara. Al final imprime un reporte ✅/❌, lo guarda en
tools/reportes-prueba/ y borra el esquema.

Uso (desde la raíz del repo):
  backend\venv\Scripts\python.exe tools\prueba_completa.py              # corrida normal
  backend\venv\Scripts\python.exe tools\prueba_completa.py --conservar  # no borra el esquema (para revisar)

Requisitos: PostgreSQL 18 (pg_dump/pg_restore en C:\Program Files\PostgreSQL\18\bin o PG_BIN_DIR) y la base
pos_test_db. Toma la conexión de TEST_DATABASE_URL o, si no existe, de backend/.env cambiando la base a
pos_test_db. Sale con código 1 si alguna prueba falla.
"""
import os
import secrets
import subprocess
import sys
import tempfile
import time
import traceback
import uuid
from datetime import datetime
from decimal import ROUND_HALF_EVEN, Decimal
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
PY = Path(sys.executable)
CENT = Decimal("0.01")


def money(value) -> Decimal:
    return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_EVEN)


# ── Base de prueba aislada ──────────────────────────────────────────────────────────────────────────


def base_url_de_prueba() -> str:
    url = os.environ.get("TEST_DATABASE_URL")
    if not url:
        from dotenv import dotenv_values

        url = dotenv_values(BACKEND / ".env").get("DATABASE_URL", "")
        parts = urlsplit(url)
        url = urlunsplit(parts._replace(path="/pos_test_db", query=""))
    parts = urlsplit(url)
    if parts.path.strip("/") != "pos_test_db":
        raise SystemExit(f"Por seguridad solo se prueba contra pos_test_db (se recibió {parts.path!r}).")
    return urlunsplit(parts._replace(query=""))


def preparar_esquema(base: str, esquema: str) -> tuple[str, dict]:
    import sqlalchemy as sa

    engine = sa.create_engine(base)
    with engine.begin() as conn:
        conn.execute(sa.text(f"CREATE SCHEMA {esquema}"))
    engine.dispose()
    url = f"{base}?options=-csearch_path%3D{esquema},public"
    env = {**os.environ, "DATABASE_URL": url, "FLASK_APP": "run.py", "POS_ADMIN_PASSWORD": secrets.token_urlsafe(12)}
    for args in (["-m", "flask", "db", "upgrade"], ["-m", "flask", "seed"]):
        r = subprocess.run([str(PY), *args], cwd=BACKEND, env=env, capture_output=True, text=True)
        if r.returncode:
            raise SystemExit(f"Falló `flask {' '.join(args[2:])}`:\n{r.stdout[-800:]}\n{r.stderr[-1500:]}")
    return url, {}


def borrar_esquema(base: str, esquema: str) -> None:
    import sqlalchemy as sa

    engine = sa.create_engine(base)
    with engine.begin() as conn:
        conn.execute(sa.text(f"DROP SCHEMA IF EXISTS {esquema} CASCADE"))
    engine.dispose()


# ── Registro de resultados ──────────────────────────────────────────────────────────────────────────


class Prueba:
    def __init__(self):
        self.resultados: list[tuple[str, str, bool, str]] = []
        self.area = ""

    def seccion(self, nombre: str):
        self.area = nombre
        print(f"\n== {nombre}")

    def check(self, nombre: str, ok: bool, detalle: str = ""):
        self.resultados.append((self.area, nombre, bool(ok), detalle))
        print(f"  {'✅' if ok else '❌'} {nombre}" + (f" — {detalle}" if detalle and not ok else ""))
        return ok

    def igual(self, nombre: str, obtenido, esperado):
        return self.check(nombre, obtenido == esperado, f"esperado {esperado!r}, obtenido {obtenido!r}")

    def status(self, nombre: str, resp, esperado):
        esperados = esperado if isinstance(esperado, (tuple, list)) else (esperado,)
        cuerpo = ""
        if resp.status_code not in esperados:
            try:
                cuerpo = str(resp.get_json())[:200]
            except Exception:  # noqa: BLE001
                cuerpo = resp.data[:200]
        return self.check(nombre, resp.status_code in esperados, f"HTTP {resp.status_code} (esperado {esperado}) {cuerpo}")

    @property
    def fallas(self):
        return [r for r in self.resultados if not r[2]]


# ── Cliente de la API ───────────────────────────────────────────────────────────────────────────────


class Api:
    def __init__(self, client):
        self.c = client
        self.token = None

    def _h(self, token):
        tok = token if token is not None else self.token
        return {"Authorization": f"Bearer {tok}"} if tok else {}

    def get(self, url, token=None):
        return self.c.get("/api" + url, headers=self._h(token))

    def post(self, url, body=None, token=None):
        return self.c.post("/api" + url, json=body if body is not None else {}, headers=self._h(token))

    def patch(self, url, body, token=None):
        return self.c.patch("/api" + url, json=body, headers=self._h(token))

    def put(self, url, body, token=None):
        return self.c.put("/api" + url, json=body, headers=self._h(token))

    def delete(self, url, token=None):
        return self.c.delete("/api" + url, headers=self._h(token))

    def login(self, usuario, contrasena):
        r = self.post("/auth/login", {"username": usuario, "password": contrasena}, token="")
        return r, (r.get_json() or {})


def items(resp):
    data = resp.get_json()
    if isinstance(data, list):
        return data
    data = data or {}
    return data.get("items", data.get("backups", []))


# ── Recorrido ───────────────────────────────────────────────────────────────────────────────────────


def recorrido(app, t: Prueba, credenciales: dict):  # noqa: C901 - es un guion lineal y legible a propósito
    api = Api(app.test_client())
    ctx: dict = {}

    # 1. Salud
    t.seccion("1. Salud del sistema")
    r = api.get("/health", token="")
    t.status("/api/health responde", r, 200)
    t.igual("base de datos ok", (r.get_json() or {}).get("db"), "ok")

    # 2. Sesión
    t.seccion("2. Inicio de sesión y sesión")
    r, _ = api.login("admin", "contraseña-incorrecta")
    t.status("contraseña incorrecta se rechaza", r, 401)
    r, data = api.login("ADMIN", credenciales["admin"])
    t.status("login del admin (sin distinguir mayúsculas)", r, 200)
    api.token = data.get("access_token")
    refresh = data.get("refresh_token")
    t.status("/auth/me", api.get("/auth/me"), 200)
    r = api.post("/auth/refresh", token=refresh)
    t.status("renovar sesión con el refresh token", r, 200)
    t.status("token alterado da 401", api.get("/products/", token="abc.def.ghi"), 401)

    # 3. Usuarios
    t.seccion("3. Usuarios y roles")
    roles = {x["name"]: x["id"] for x in items(api.get("/users/roles"))}
    t.check("roles admin/gerente/cajero existen", {"admin", "gerente", "cajero"} <= set(roles))
    pw_cajero = "Prueba-" + secrets.token_urlsafe(10)
    r = api.post("/users/", {"username": "  cajero_prueba  ", "full_name": "Cajero Prueba", "password": pw_cajero,
                             "role_id": roles["cajero"], "email": "cajero.prueba@example.com"})
    t.status("crear cajero (recorta espacios)", r, 201)
    cajero_id = (r.get_json() or {}).get("id")
    t.igual("username guardado sin espacios", (r.get_json() or {}).get("username"), "cajero_prueba")
    r = api.post("/users/", {"username": "CAJERO_PRUEBA", "full_name": "Duplicado", "password": pw_cajero, "role_id": roles["cajero"]})
    t.status("username duplicado en mayúsculas se rechaza", r, (400, 409, 422))
    r = api.post(f"/users/{cajero_id}/reset-password")
    t.status("restablecer contraseña", r, 200)
    temporal = (r.get_json() or {}).get("temporary_password", "")
    t.check("contraseña temporal de 12+ caracteres", len(temporal) >= 12)
    r, data = api.login("cajero_prueba", temporal)
    t.status("login con la temporal", r, 200)
    t.igual("obliga a cambiar la contraseña", data.get("must_change_password") or (data.get("user") or {}).get("must_change_password"), True)
    tok_tmp = data.get("access_token")
    t.status("API bloqueada hasta cambiarla", api.get("/products/", token=tok_tmp), 403)
    r = api.post("/auth/change-password", {"current_password": temporal, "new_password": pw_cajero}, token=tok_tmp)
    t.status("cambiar la contraseña", r, 200)
    t.check("el cambio entrega un refresh token nuevo", bool((r.get_json() or {}).get("refresh_token")))
    r, data = api.login("cajero_prueba", pw_cajero)
    t.status("login del cajero con la nueva", r, 200)
    ctx["cajero"] = data.get("access_token")

    # 4. Permisos por rol
    t.seccion("4. Permisos por rol (cajero)")
    for url in ("/users/", "/config/", "/reports/dashboard", "/audit/", "/backup/", "/purchases/orders"):
        t.status(f"cajero no entra a {url}", api.get(url, token=ctx["cajero"]), 403)
    t.status("sin sesión no hay acceso", api.get("/products/", token=""), 401)

    # 5. Catálogo
    t.seccion("5. Catálogo: unidades, categorías y productos")
    unidades = {u["abbreviation"]: u for u in items(api.get("/products/units"))}
    if "pz" not in unidades:
        unidades["pz"] = api.post("/products/units", {"name": "Pieza", "abbreviation": "pz", "allows_fractions": False}).get_json()
    r = api.post("/products/units", {"name": "Kilogramo prueba", "abbreviation": "kgp", "allows_fractions": True})
    t.status("crear unidad a granel", r, 201)
    unidades["kgp"] = r.get_json()
    t.status("unidad duplicada se rechaza", api.post("/products/units", {"name": "KILOGRAMO PRUEBA ", "abbreviation": "kgx"}), (400, 409, 422))
    r = api.post("/categories/", {"name": "Herramientas prueba"})
    t.status("crear categoría", r, 201)
    cat = (r.get_json() or {}).get("id")
    t.status("categoría en blanco se rechaza", api.post("/categories/", {"name": "   "}), (400, 422))

    def alta(sku, barcode, nombre, unidad, precio, costo):
        return api.post("/products/", {"sku": sku, "barcode": barcode, "name": nombre, "category_id": cat,
                                       "unit_id": unidades[unidad]["id"], "sale_price": precio, "purchase_price": costo,
                                       "tax_rate": "16.00", "min_stock": "5"})

    r = alta("PRB-MAR", "7509990000011", "Martillo prueba", "pz", "100.00", "60.00")
    t.status("crear producto por pieza", r, 201)
    martillo = r.get_json() or {}
    r = alta("PRB-CLV", "7509990000028", "Clavo a granel prueba", "kgp", "40.00", "25.00")
    t.status("crear producto a granel", r, 201)
    clavo = r.get_json() or {}
    t.igual("costo promedio inicial = precio de compra", money(martillo.get("avg_cost", 0)), Decimal("60.00"))
    t.status("SKU duplicado (mayúsculas/espacios) se rechaza", alta(" prb-mar ", None, "dup", "pz", "1", "1"), (400, 409, 422))
    t.status("precio con 3 decimales se rechaza", alta("PRB-3D", None, "x", "pz", "10.125", "1"), 422)
    r = api.get("/products/by-code/7509990000011")
    t.igual("escaneo por código de barras", (r.get_json() or {}).get("sku"), "PRB-MAR")
    t.igual("escaneo por SKU", (api.get("/products/by-code/prb-clv").get_json() or {}).get("sku"), "PRB-CLV")
    r = api.get("/products/?search=martilo&page=1&per_page=20")
    t.check("búsqueda tolera errores de dedo", any(p["sku"] == "PRB-MAR" for p in items(r)))
    t.check("buscar % no devuelve todo", len(items(api.get("/products/?search=%25&page=1&per_page=50"))) == 0)
    r = api.get("/products/pos-catalog")
    t.check("catálogo del POS incluye los productos", {"PRB-MAR", "PRB-CLV"} <= {p["sku"] for p in items(r)})
    r = api.post(f"/printing/products/{martillo['id']}/label", {"preview": True})
    t.check("etiqueta con precio con IVA ($116.00)", "$116.00" in str(r.get_json()), f"HTTP {r.status_code}")

    # 6. Compras
    t.seccion("6. Compras: proveedor y orden con recepción parcial")
    r = api.post("/purchases/suppliers", {"name": "Proveedor Prueba", "rfc": "PRU010101AB1", "phone": "4771234567"})
    t.status("crear proveedor", r, 201)
    prov = (r.get_json() or {}).get("id")
    t.status("RFC con formato inválido se rechaza", api.post("/purchases/suppliers", {"name": "Otro", "rfc": "123"}), (400, 422))
    r = api.post("/purchases/orders", {"supplier_id": prov, "items": [
        {"product_id": martillo["id"], "quantity": "20", "unit_cost": "60.00"},
        {"product_id": clavo["id"], "quantity": "50", "unit_cost": "25.00"}]})
    t.status("crear orden de compra", r, 201)
    orden = r.get_json() or {}
    t.check("folio de OC", str(orden.get("folio", "")).startswith("OC-"))
    t.status("enviar orden", api.post(f"/purchases/orders/{orden['id']}/send"), 200)
    oi = {i["product_id"]: i["id"] for i in orden.get("items", [])}
    r = api.post(f"/purchases/orders/{orden['id']}/receive", {"items": [
        {"order_item_id": oi[martillo["id"]], "received_quantity": "20"},
        {"order_item_id": oi[clavo["id"]], "received_quantity": "30"}]})
    t.status("recepción parcial", r, 200)
    t.igual("orden queda parcial", (api.get(f"/purchases/orders/{orden['id']}").get_json() or {}).get("status"), "parcial")
    t.status("recibir más de lo pedido se rechaza", api.post(f"/purchases/orders/{orden['id']}/receive", {"items": [
        {"order_item_id": oi[clavo["id"]], "received_quantity": "999"}]}), (400, 422))
    r = api.post(f"/purchases/orders/{orden['id']}/close", {"notes": "Proveedor sin existencia"})
    t.status("cerrar orden parcial", r, 200)
    t.igual("orden queda cerrada", (api.get(f"/purchases/orders/{orden['id']}").get_json() or {}).get("status"), "cerrada")

    # 7. Inventario
    t.seccion("7. Inventario: existencias, ajustes y conteo físico")

    def stock(pid):
        return Decimal(str((api.get(f"/products/{pid}").get_json() or {}).get("stock_qty", "0")))

    t.igual("stock del martillo tras la compra", stock(martillo["id"]), Decimal("20"))
    t.igual("stock del clavo tras la compra", stock(clavo["id"]), Decimal("30"))
    r = api.post("/inventory/adjust", {"product_id": martillo["id"], "quantity": "22", "reason": "error_captura", "notes": "prueba"})
    t.status("ajuste con motivo", r, (200, 201))
    t.igual("stock ajustado", stock(martillo["id"]), Decimal("22"))
    t.status("ajuste fraccionario de pieza se rechaza", api.post("/inventory/adjust", {"product_id": martillo["id"], "quantity": "2.5", "reason": "otro"}), (400, 422))
    t.status("motivo inválido se rechaza", api.post("/inventory/adjust", {"product_id": martillo["id"], "quantity": "22", "reason": "porque_si"}), (400, 422))
    movs = items(api.get(f"/inventory/movements/{martillo['id']}"))
    t.check("historial de movimientos", len(movs) >= 2)
    r = api.get("/inventory/summary")
    t.status("resumen de inventario (KPI)", r, 200)
    r = api.post("/inventory/counts", {"notes": "Conteo de prueba"})
    t.status("crear conteo físico", r, (200, 201))
    conteo = r.get_json() or {}
    r = api.patch(f"/inventory/counts/{conteo.get('id')}/items/{martillo['id']}", {"counted_quantity": "21", "notes": "faltó 1"})
    t.status("capturar conteo", r, 200)
    r = api.post(f"/inventory/counts/{conteo.get('id')}/close")
    t.status("cerrar conteo aplicando la diferencia", r, 200)
    t.igual("stock tras el conteo", stock(martillo["id"]), Decimal("21"))
    t.status("conteo inexistente da 404", api.get("/inventory/counts/999999"), 404)

    # 7b. Inventario con Excel
    t.seccion("7b. Inventario con Excel (exportar, vista previa y aplicar)")
    from io import BytesIO

    from openpyxl import load_workbook

    r = api.get("/inventory/export")
    t.check("exportar a Excel", r.status_code == 200 and r.data[:2] == b"PK", f"HTTP {r.status_code}")
    wb = load_workbook(BytesIO(r.data))
    r.close()
    ws = wb["Inventario"]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}

    def subir(url, libro, **form):
        buf = BytesIO()
        libro.save(buf)
        buf.seek(0)
        return api.c.post("/api" + url, headers=api._h(None), content_type="multipart/form-data",
                          data={"archivo": (buf, "inventario.xlsx"), **form})

    r = subir("/inventory/import/preview", wb)
    t.igual("reimportar lo exportado: sin cambios", (r.get_json() or {}).get("resumen", {}).get("cambios"), 0)
    for fila in range(2, ws.max_row + 1):
        if ws.cell(fila, col["SKU"]).value == "PRB-CLV":
            ws.cell(fila, col["Existencias"]).value = 29.75
    nueva = ws.max_row + 1
    for k, v in {"SKU": "PRB-NIV", "Nombre": "Nivel prueba", "Unidad": "pz", "Precio de venta": 145, "Existencias": 3}.items():
        ws.cell(nueva, col[k]).value = v
    vista = subir("/inventory/import/preview", wb).get_json() or {}
    t.igual("vista previa: 1 nuevo y 1 con cambios", (vista.get("resumen", {}).get("nuevos"), vista.get("resumen", {}).get("cambios")), (1, 1))
    r = subir("/inventory/import/apply", wb, plan_hash=vista.get("plan_hash", ""))
    t.status("aplicar la importación", r, 200)
    t.igual("existencias reemplazadas desde el Excel", stock(clavo["id"]), Decimal("29.75"))
    t.check("producto nuevo creado desde el Excel", (api.get("/products/by-code/PRB-NIV").get_json() or {}).get("sku") == "PRB-NIV")
    ajuste = api.post("/inventory/adjust", {"product_id": clavo["id"], "quantity": "30", "reason": "otro"})
    t.status("devolver el clavo a 30 kg para el resto del recorrido", ajuste, (200, 201))

    # 8. Caja
    t.seccion("8. Cajas y apertura de turno")
    r = api.post("/cash-register/registers", {"name": "Caja prueba"})
    t.status("dar de alta una caja", r, 201)
    caja = (r.get_json() or {}).get("id")
    t.status("nombre de caja repetido se rechaza", api.post("/cash-register/registers", {"name": "CAJA PRUEBA"}), 400)
    r = api.post("/cash-register/open", {"cash_register_id": caja, "opening_amount": "500.00"}, token=ctx["cajero"])
    t.status("el cajero abre turno", r, 201)
    sesion = (r.get_json() or {}).get("id")
    t.status("abrir la misma caja otra vez se rechaza", api.post("/cash-register/open", {"cash_register_id": caja, "opening_amount": "1"}), (400, 409))

    # 9. Ventas
    t.seccion("9. Ventas")
    cajero = ctx["cajero"]

    def vender(lineas, pagos, client_uuid=None, token=None):
        return api.post("/sales/", {"session_id": sesion, "client_uuid": client_uuid or str(uuid.uuid4()),
                                    "items": lineas, "payments": pagos}, token=token or cajero)

    r = api.post("/sales/quote", {"session_id": sesion, "items": [{"product_id": martillo["id"], "quantity": "2"}]}, token=cajero)
    t.status("cotización", r, 200)
    t.igual("cotización: 2 × 100 + 16 % = 232.00", money((r.get_json() or {}).get("total", 0)), Decimal("232.00"))
    r = vender([{"product_id": martillo["id"], "quantity": "2", "unit_price": "1.00"}], [{"method": "efectivo", "amount": "300.00"}])
    t.status("venta en efectivo", r, 201)
    v1 = r.get_json() or {}
    t.igual("el servidor fija el precio (ignora el del cliente)", money(v1.get("total", 0)), Decimal("232.00"))
    t.igual("cambio calculado por el servidor", money(sum(Decimal(p["change_amount"]) for p in v1.get("payments", []))), Decimal("68.00"))
    t.check("folio de venta", str(v1.get("folio", "")).startswith("V-"))
    r = vender([{"product_id": clavo["id"], "quantity": "2.5"}], [{"method": "efectivo", "amount": "50.00"},
                                                                   {"method": "tarjeta", "amount": "66.00", "reference": "4455"}])
    t.status("venta a granel con pago mixto", r, 201)
    v2 = r.get_json() or {}
    t.igual("2.5 kg × 40 + 16 % = 116.00", money(v2.get("total", 0)), Decimal("116.00"))
    t.status("cantidad con más de 3 decimales se rechaza", vender([{"product_id": clavo["id"], "quantity": "1.0004"}], [{"method": "efectivo", "amount": "100"}]), 422)
    t.status("media pieza se rechaza", vender([{"product_id": martillo["id"], "quantity": "0.5"}], [{"method": "efectivo", "amount": "100"}]), (400, 422))
    t.status("stock insuficiente se rechaza", vender([{"product_id": martillo["id"], "quantity": "999"}], [{"method": "efectivo", "amount": "999999"}]), (400, 409, 422))
    t.status("pago insuficiente se rechaza", vender([{"product_id": martillo["id"], "quantity": "1"}], [{"method": "efectivo", "amount": "10.00"}]), (400, 422))
    uid = str(uuid.uuid4())
    r1 = vender([{"product_id": martillo["id"], "quantity": "1"}], [{"method": "efectivo", "amount": "116.00"}], client_uuid=uid)
    r2 = vender([{"product_id": martillo["id"], "quantity": "1"}], [{"method": "efectivo", "amount": "116.00"}], client_uuid=uid)
    t.check("reenvío idempotente (misma venta, sin duplicar)", r1.status_code == 201 and (r2.get_json() or {}).get("id") == (r1.get_json() or {}).get("id"),
            f"{r1.status_code}/{r2.status_code}")
    v3 = r1.get_json() or {}
    t.igual("stock del martillo: 21 − 2 − 1 = 18", stock(martillo["id"]), Decimal("18"))
    t.igual("stock del clavo: 30 − 2.5 = 27.5", stock(clavo["id"]), Decimal("27.5"))
    r = api.post(f"/printing/sales/{v1.get('id')}/ticket", {"preview": True}, token=cajero)
    t.check("vista previa del ticket", v1.get("folio", "") in str(r.get_json()), f"HTTP {r.status_code}")
    lista = items(api.get(f"/sales/?page=1&per_page=20&user_id={cajero_id}"))
    t.check("historial filtrado por cajero", len(lista) >= 3)

    # 10. Devoluciones y cancelación
    t.seccion("10. Devoluciones y cancelación")

    def devolver(venta, qty, restock=True, item=0):
        return api.post(f"/sales/{venta['id']}/return", {"reason": "Prueba", "session_id": sesion, "refund_method": "efectivo",
                                                          "restock": restock, "items": [{"sale_item_id": venta["items"][item]["id"], "quantity": str(qty)}]})

    t.status("devolver media pieza se rechaza", devolver(v1, "0.5"), (400, 422))
    r1 = devolver(v1, 1)
    t.status("devolución parcial", r1, 200)
    t.check("folio de devolución D-", str((r1.get_json() or {}).get("folio", "")).startswith("D-"))
    r2 = devolver(v1, 1, restock=False)
    t.status("devolución sin regresar al inventario (merma)", r2, 200)
    reemb = money((r1.get_json() or {}).get("refund_total", 0)) + money((r2.get_json() or {}).get("refund_total", 0))
    t.igual("reembolsos suman exactamente lo cobrado", reemb, money(v1["total"]))
    t.igual("stock: +1 con reposición, +0 merma → 19", stock(martillo["id"]), Decimal("19"))
    t.status("devolver más de lo vendido se rechaza", devolver(v1, 1), (400, 422))
    r = api.post(f"/sales/{v3.get('id')}/cancel", {"reason": "Cliente se arrepintió"})
    t.status("cancelar venta", r, 200)
    t.igual("cancelar regresa el stock → 20", stock(martillo["id"]), Decimal("20"))
    t.status("cancelar sin motivo se rechaza", api.post(f"/sales/{v2.get('id')}/cancel", {}), (400, 422))
    t.status("venta inexistente da 404", api.post("/sales/999999/cancel", {"reason": "x"}), 404)

    # 11. Corte de caja
    t.seccion("11. Corte de caja")
    # Efectivo esperado = fondo + efectivo neto de ventas vigentes − reembolsos en efectivo de esta caja
    efectivo_v1 = sum(Decimal(p["amount"]) - Decimal(p["change_amount"]) for p in v1["payments"] if p["method"] == "efectivo")
    efectivo_v2 = sum(Decimal(p["amount"]) - Decimal(p["change_amount"]) for p in v2["payments"] if p["method"] == "efectivo")
    esperado = money(Decimal("500.00") + efectivo_v1 + efectivo_v2 - reemb)
    r = api.get(f"/cash-register/{sesion}/summary", token=cajero)
    t.status("resumen del turno", r, 200)
    r = api.post(f"/cash-register/{sesion}/close", {"closing_amount": str(esperado - Decimal("5.00")), "close_notes": "Faltante de prueba"}, token=cajero)
    t.status("cerrar turno", r, 200)
    cierre = r.get_json() or {}
    t.igual("efectivo esperado (recalculado a mano)", money(cierre.get("expected_cash", 0)), esperado)
    t.igual("diferencia −5.00", money(cierre.get("difference", 0)), Decimal("-5.00"))
    t.igual("observaciones guardadas", cierre.get("close_notes"), "Faltante de prueba")

    # 12. Reportes
    t.seccion("12. Reportes")
    hoy = datetime.now().strftime("%Y-%m-%d")
    for url in ("/reports/dashboard", f"/reports/sales?from={hoy}&to={hoy}", "/reports/top-products?limit=10",
                "/reports/inventory-valuation", f"/reports/gross-margin?from={hoy}&to={hoy}"):
        t.status(f"GET {url.split('?')[0]}", api.get(url), 200)
    ventas = api.get(f"/reports/sales?from={hoy}&to={hoy}").get_json() or {}
    t.check("el reporte de ventas cuenta las ventas de hoy", (ventas.get("total_sales") or 0) >= 2)
    t.status("rango invertido se rechaza", api.get(f"/reports/sales?from={hoy}&to=2000-01-01"), 400)

    # 13. Configuración y auditoría
    t.seccion("13. Configuración y auditoría")
    r = api.put("/config/", {"company_name": "Ferretería de Prueba", "max_discount_percent": "10"})
    t.status("guardar configuración válida", r, 200)
    t.status("tope de descuento inválido se rechaza", api.put("/config/", {"max_discount_percent": "150"}), 400)
    t.igual("branding con el nombre nuevo", (api.get("/config/branding", token="").get_json() or {}).get("company_name"), "Ferretería de Prueba")
    auditoria = items(api.get("/audit/?page=1&per_page=100"))
    acciones = {a.get("action") for a in auditoria}
    t.check("la auditoría registra ventas, devoluciones y ajustes", len(auditoria) > 10 and bool(acciones))
    texto = str(auditoria)
    t.check("la auditoría no guarda contraseñas", pw_cajero not in texto and temporal not in texto)

    # 14. Respaldos
    t.seccion("14. Respaldos")
    r = api.post("/backup/")
    t.status("crear respaldo", r, 201)
    nombre = (r.get_json() or {}).get("filename")
    t.check("respaldo aparece en la lista", any(b.get("filename") == nombre for b in items(api.get("/backup/"))))
    t.status("verificar respaldo", api.post(f"/backup/{nombre}/verify"), 200)
    r = api.get(f"/backup/{nombre}/download")
    t.check("descargar respaldo", r.status_code == 200 and len(r.data) > 1000, f"HTTP {r.status_code}, {len(r.data)} bytes")
    r.close()  # en Windows un archivo abierto no se puede borrar: el servidor real cierra la respuesta al enviarla
    t.status("nombre de archivo malicioso se rechaza", api.delete("/backup/..%2F..%2Fwin.ini"), (400, 404))
    t.status("borrar respaldo", api.delete(f"/backup/{nombre}"), 200)

    # 15. Cierre de sesión
    t.seccion("15. Cierre de sesión")
    t.status("logout", api.post("/auth/logout"), 200)
    t.status("el token ya no sirve", api.get("/auth/me"), 401)


def main() -> int:
    conservar = "--conservar" in sys.argv
    base = base_url_de_prueba()
    esquema = "prueba_" + datetime.now().strftime("%Y%m%d_%H%M%S")
    print(f"Base de prueba: pos_test_db, esquema nuevo {esquema}")
    inicio = time.time()
    url, _ = preparar_esquema(base, esquema)

    respaldos = tempfile.mkdtemp(prefix="pos-prueba-respaldos-")
    os.environ.update({"DATABASE_URL": url, "BACKUP_DIR": respaldos, "USB_BACKUP_DIR": ""})
    os.environ.setdefault("PG_BIN_DIR", r"C:\Program Files\PostgreSQL\18\bin")
    sys.path.insert(0, str(BACKEND))
    os.chdir(BACKEND)

    from app import create_app
    from app.extensions import db
    from app.models.user import Role, User

    app = create_app()
    app.config.update(BACKUP_PG_SCHEMA=esquema, RESTORE_TERMINATE_CONNECTIONS=False, USB_BACKUP_DIR=None,
                      MAINTENANCE_FLAG_FILE=os.path.join(respaldos, "mantenimiento.flag"))
    t = Prueba()
    try:
        with app.app_context():
            actual = db.session.execute(db.text("SELECT current_database() || '/' || current_schema()")).scalar()
            if actual != f"pos_test_db/{esquema}":
                raise SystemExit(f"Conexión inesperada ({actual}); se aborta por seguridad.")
            admin = User.query.join(Role).filter(Role.name == "admin").first()
            credenciales = {"admin": "Prueba-" + secrets.token_urlsafe(12)}
            admin.set_password(credenciales["admin"])
            admin.must_change_password = False
            db.session.commit()
        try:
            recorrido(app, t, credenciales)
        except Exception:  # noqa: BLE001 - un error inesperado se reporta como falla, no detiene el reporte
            t.check("el recorrido terminó sin errores inesperados", False, traceback.format_exc(limit=3))
    finally:
        if not conservar:
            borrar_esquema(base, esquema)

    total, fallas = len(t.resultados), t.fallas
    duracion = time.time() - inicio
    print(f"\nResultado: {total - len(fallas)}/{total} pruebas pasaron en {duracion:.0f} s"
          + ("" if not fallas else f" — {len(fallas)} FALLARON"))
    for area, nombre, _ok, detalle in fallas:
        print(f"  ❌ [{area}] {nombre}: {detalle}")

    salida = ROOT / "tools" / "reportes-prueba"
    salida.mkdir(exist_ok=True)
    archivo = salida / f"prueba-{datetime.now():%Y%m%d-%H%M%S}.md"
    lineas = [f"# Prueba completa del POS — {datetime.now():%Y-%m-%d %H:%M}", "",
              f"**Resultado:** {total - len(fallas)}/{total} pasaron ({duracion:.0f} s). Esquema: `{esquema}`"
              + (" (conservado)" if conservar else " (borrado)"), ""]
    area_actual = None
    for area, nombre, ok, detalle in t.resultados:
        if area != area_actual:
            lineas += ["", f"## {area}", ""]
            area_actual = area
        lineas.append(f"- {'✅' if ok else '❌'} {nombre}" + (f" — {detalle}" if not ok and detalle else ""))
    archivo.write_text("\n".join(lineas) + "\n", encoding="utf-8")
    print(f"Reporte: {archivo.relative_to(ROOT)}")
    return 1 if fallas else 0


if __name__ == "__main__":
    sys.exit(main())
